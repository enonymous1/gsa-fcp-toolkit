"""Presentation helpers for the CPL Word documents (python-docx)."""

from __future__ import annotations

from docx.document import Document as DocxDocument
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from gsa_fcp_toolkit.config import (
    COLOR_PLACEHOLDER_RED,
    COLOR_SLATE_GRAY,
    COLOR_WHITE,
    CPL_TABLE_COLUMN_WIDTHS,
)


def rgb(color: tuple[int, int, int]) -> RGBColor:
    """Convert an (R, G, B) tuple from config to a docx RGBColor."""
    return RGBColor(*color)


def _hex(color: tuple[int, int, int]) -> str:
    """Convert an (R, G, B) tuple to a hex string like '4F6272'."""
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_column_widths(table: Table, widths: tuple[float, ...] = CPL_TABLE_COLUMN_WIDTHS) -> None:
    """Set column widths (python-docx requires widths on every cell)."""
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def apply_header_styling(cell: _Cell) -> None:
    """Slate-gray background with white bold centered text for header cells."""
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run.font.bold = True
        run.font.color.rgb = rgb(COLOR_WHITE)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{_hex(COLOR_SLATE_GRAY)}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def apply_placeholder_styling(run: Run) -> None:
    """Style a run red and bold to flag required manual input."""
    run.font.color.rgb = rgb(COLOR_PLACEHOLDER_RED)
    run.font.bold = True


def add_paragraph_with_placeholders(doc: DocxDocument, text: str) -> Paragraph:
    """Add a paragraph, styling [BRACKETED] segments red and bold."""
    para = doc.add_paragraph()
    parts = text.split("[")

    if parts[0]:
        para.add_run(parts[0])

    for part in parts[1:]:
        if "]" in part:
            bracket_text, normal_text = part.split("]", 1)
            placeholder_run = para.add_run(f"[{bracket_text}]")
            apply_placeholder_styling(placeholder_run)
            if normal_text:
                para.add_run(normal_text)
        else:
            para.add_run(f"[{part}")  # malformed bracket fallback

    return para
