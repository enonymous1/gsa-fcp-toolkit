"""Assemble the Commercial Price List Word documents.

Pure document assembly: data loading, cleaning, and sorting happen upstream
(the CLI pipeline); this module only renders ``labor_data`` to .docx.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from gsa_fcp_toolkit import config
from gsa_fcp_toolkit.models import LaborItem
from gsa_fcp_toolkit.pricing import calculate_commercial_rate, calculate_gsa_rate
from gsa_fcp_toolkit.word.styles import (
    add_paragraph_with_placeholders,
    apply_header_styling,
    rgb,
    set_column_widths,
)

logger = logging.getLogger(__name__)

DOCUMENT_TITLES = {
    "proposal": "Established Commercial Price List",
    "fss": "Authorized FSS Pricelist",
}
RATE_COLUMN_TITLES = {
    "proposal": "Commercial Hourly Rate",
    "fss": "GSA Hourly Rate (w/ IFF)",
}


def build_proposal_header(doc: DocxDocument) -> None:
    """Corporate/terms header for the Established Pricing document."""
    doc.add_paragraph()  # spacing
    add_paragraph_with_placeholders(
        doc, f"Corporate Address: {config.COMPANY_LOCATION} [ENTER FULL STREET ADDRESS]"
    )
    add_paragraph_with_placeholders(
        doc, f"Point of Contact: {config.COMPANY_POC} (or [ENTER PREFERRED ADMINISTRATIVE POC])"
    )
    add_paragraph_with_placeholders(doc, "Unique Entity Identifier (UEI): [ENTER UEI NUMBER]")
    doc.add_paragraph()

    doc.add_heading("Standard Commercial Terms", level=2)
    doc.add_paragraph(f"Payment Terms: {config.PAYMENT_TERMS}")
    add_paragraph_with_placeholders(doc, "Volume Discounts: [ENTER APPLICABLE DISCOUNTS OR 'NONE']")
    doc.add_paragraph()


def build_fss_header(doc: DocxDocument) -> None:
    """Mandatory Customer Information block for the GSA FSS pricelist."""
    doc.add_paragraph()  # spacing
    doc.add_heading("Customer Information (I-FSS-600)", level=2)

    add_paragraph_with_placeholders(doc, "Schedule Title: Multiple Award Schedule (MAS)")
    add_paragraph_with_placeholders(doc, "FSC Group/Class: [ENTER FSC GROUP/CLASS]")
    add_paragraph_with_placeholders(doc, "Contract Number: [ENTER CONTRACT NUMBER]")
    add_paragraph_with_placeholders(doc, "Period of Performance: [ENTER PERIOD OF PERFORMANCE]")
    add_paragraph_with_placeholders(doc, "Maximum Order: [ENTER MAXIMUM ORDER THRESHOLD]")
    add_paragraph_with_placeholders(doc, "Minimum Order: [ENTER MINIMUM ORDER THRESHOLD]")
    doc.add_paragraph(
        "Geographic Coverage (Delivery Area): Domestic, 50 states, Washington, DC, "
        "Puerto Rico, US Territories"
    )
    doc.add_paragraph(f"Points of Production: {config.COMPANY_LOCATION}, US")
    add_paragraph_with_placeholders(
        doc,
        "Prompt Payment Terms: [ENTER PROMPT PAYMENT TERMS, e.g., Net 30. Information for "
        "Ordering Offices: Prompt payment terms cannot be negotiated out of the requisite document.]",
    )
    add_paragraph_with_placeholders(
        doc, "Delivery Terms: [ENTER DELIVERY TERMS, e.g., specified on the Task Order]"
    )
    doc.add_paragraph()


def _build_title_block(doc: DocxDocument, document_title: str) -> None:
    """Centered company name, document title, and effective date."""
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(config.COMPANY_NAME)
    title_run.font.name = config.DOC_FONT_NAME
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = rgb(config.COLOR_CORPORATE_BLUE)

    subtitle_para = doc.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_para.add_run(document_title)
    subtitle_run.font.name = config.DOC_FONT_NAME
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.bold = True

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_para.add_run(f"Effective Date: {datetime.now():%B %d, %Y}")
    date_run.font.name = config.DOC_FONT_NAME
    date_run.font.size = Pt(12)
    date_run.font.italic = True


def _build_rates_table(doc: DocxDocument, document_type: str, labor_data: list[LaborItem]) -> None:
    """Append the labor category / description / rate table."""
    logger.info("Building rates table for %d items...", len(labor_data))
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Labor Category & Qualifications"
    hdr_cells[1].text = "Description"
    hdr_cells[2].text = RATE_COLUMN_TITLES[document_type]
    for cell in hdr_cells:
        apply_header_styling(cell)
    set_column_widths(table)

    for item in labor_data:
        title = item.get("title", "Unknown Role")
        desc = item.get("description", "No description provided.")
        if document_type == "proposal":
            calc_rate = calculate_commercial_rate(item)
        else:
            calc_rate = calculate_gsa_rate(item)
        rate_str = f"${calc_rate:,.2f}"

        row_cells = table.add_row().cells

        cat_para = row_cells[0].paragraphs[0]
        cat_run = cat_para.add_run(title)
        cat_run.font.bold = True

        min_edu = item.get("minimum_education", "N/A")
        min_exp = item.get("minimum_years_of_experience", "0")
        qual_run = cat_para.add_run(f"\nEducation: {min_edu}\nExperience: {min_exp} years")
        qual_run.font.size = Pt(9)
        qual_run.font.italic = True

        row_cells[1].text = desc

        rate_para = row_cells[2].paragraphs[0]
        rate_para.add_run(rate_str)
        rate_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def generate_word_cpl(document_type: str, labor_data: list[LaborItem]) -> Path:
    """Generate the Word CPL document and return the output path.

    Args:
        document_type: ``'proposal'`` for commercial rates, ``'fss'`` for
            GSA authorized (IFF-inclusive) rates.
        labor_data: Cleaned and sorted labor records (prepared upstream).

    Raises:
        ValueError: If ``document_type`` is not ``'proposal'`` or ``'fss'``.
    """
    if document_type not in ("proposal", "fss"):
        raise ValueError("Invalid document_type. Must be 'proposal' or 'fss'")
    logger.info("Generating '%s' Word document for %d items...", document_type, len(labor_data))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(config.DOC_MARGIN_INCHES)
        section.bottom_margin = Inches(config.DOC_MARGIN_INCHES)
        section.left_margin = Inches(config.DOC_MARGIN_INCHES)
        section.right_margin = Inches(config.DOC_MARGIN_INCHES)

    _build_title_block(doc, DOCUMENT_TITLES[document_type])

    if document_type == "proposal":
        build_proposal_header(doc)
    else:
        build_fss_header(doc)

    intro_txt = (
        "The following document outlines the hourly rates for professional services "
        f"provided by {config.COMPANY_NAME}."
    )
    if document_type == "proposal":
        intro_txt += " These rates reflect our standard commercial market pricing."
    else:
        intro_txt += (
            " These rates are authorized GSA schedule rates "
            "inclusive of the Industrial Funding Fee (IFF)."
        )
    intro_para = doc.add_paragraph(intro_txt)
    intro_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()  # spacing

    _build_rates_table(doc, document_type, labor_data)

    filename = (
        config.PROPOSAL_CPL_FILENAME
        if document_type == "proposal"
        else config.FSS_PRICELIST_FILENAME
    )
    output_path = config.OUTPUT_DIR / filename
    output_path.parent.mkdir(parents=True, exist_ok=True)

    logger.info("Saving %s document to %s", document_type, output_path)
    doc.save(str(output_path))
    logger.info("SUCCESS: Document generated: %s", output_path.name)
    return output_path
