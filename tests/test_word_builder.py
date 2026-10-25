"""Tests for gsa_fcp_toolkit.word.builder (python-docx runs anywhere)."""

import pytest
from docx import Document

from gsa_fcp_toolkit import config
from gsa_fcp_toolkit.word import generate_word_cpl

ITEMS = [
    {
        "title": "Engineer",
        "description": "Builds things.",
        "minimum_education": "BS",
        "minimum_years_of_experience": 3,
        "commercial_price": "100.00",
        "price_offered_to_gsa_including_iff": "95.00",
    },
]


def test_generate_proposal_cpl(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    out = generate_word_cpl("proposal", ITEMS)
    assert out == tmp_path / "Millennium_Partners_Proposal_CPL.docx"
    assert out.is_file()

    doc = Document(str(out))
    table = doc.tables[0]
    assert len(table.rows) == 2  # header + one data row
    assert table.rows[0].cells[2].text == "Commercial Hourly Rate"
    assert table.rows[1].cells[2].text == "$100.00"


def test_generate_fss_cpl_uses_gsa_rate(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    out = generate_word_cpl("fss", ITEMS)
    assert out == tmp_path / "Millennium_Partners_Authorized_FSS_Pricelist.docx"

    doc = Document(str(out))
    table = doc.tables[0]
    assert table.rows[0].cells[2].text == "GSA Hourly Rate (w/ IFF)"
    assert table.rows[1].cells[2].text == "$95.00"


def test_generate_word_cpl_rejects_bad_type(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    with pytest.raises(ValueError):
        generate_word_cpl("bogus", ITEMS)
