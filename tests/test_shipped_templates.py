"""The templates shipped in templates/ stay in sync with the version policy.

Guards against silently replacing a shipped template with a newer GSA release
without updating SUPPORTED_TEMPLATE_VERSIONS, or with a file whose Pricing
layout drifted from the other release.
"""

from pathlib import Path

import openpyxl

from gsa_fcp_toolkit.config import SUPPORTED_TEMPLATE_VERSIONS

TEMPLATES_DIR = Path(__file__).parent.parent / "templates"
EXPECTED_FILES = {
    "FCP_Services_Plus_File.xlsx": "1.1-20260602",
    "FCP_Services_Plus_File_v-1.1-20260128.xlsx": "1.1-20260128",
}


def _pricing_headers(name):
    wb = openpyxl.load_workbook(TEMPLATES_DIR / name, data_only=True, read_only=True)
    row = next(wb["Pricing"].iter_rows(min_row=2, max_row=2, max_col=100, values_only=True))
    headers = tuple(row)
    while headers and headers[-1] is None:
        headers = headers[:-1]
    return headers


def test_shipped_templates_exist():
    for name in EXPECTED_FILES:
        assert (TEMPLATES_DIR / name).is_file()


def test_shipped_template_versions_match_policy():
    found = []
    for name, expected in EXPECTED_FILES.items():
        wb = openpyxl.load_workbook(TEMPLATES_DIR / name, data_only=True, read_only=True)
        version = wb["Read Me"]["B2"].value
        assert version == expected
        found.append(version)
    assert sorted(found) == sorted(SUPPORTED_TEMPLATE_VERSIONS)


def test_shipped_templates_share_pricing_layout():
    layouts = [_pricing_headers(name) for name in EXPECTED_FILES]
    assert len(layouts[0]) == 60
    assert layouts[0] == layouts[1]
